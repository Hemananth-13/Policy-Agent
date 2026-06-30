import os
import re
import uuid
import logging
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from models import TaskItem, ReflectionNote, DocumentError

logger = logging.getLogger(__name__)

OUTPUT_DIR = Path(__file__).parent / "generated_docs"
OUTPUT_DIR.mkdir(exist_ok=True)

INSURER_NAME    = "SecureHealth Insurance Co."
INSURER_ADDRESS = "123 Main Street, Anytown, USA 12345"
INSURER_PHONE   = "(800) 555-0100"
INSURER_WEB     = "www.securehealthins.com"

DARK_BLUE = RGBColor(0x1F, 0x49, 0x7D)
MID_BLUE  = RGBColor(0x2E, 0x74, 0xB5)
GREY      = RGBColor(0x80, 0x80, 0x80)
BLACK     = RGBColor(0x00, 0x00, 0x00)


# ── Helpers ────────────────────────────────────────────────────────────────────

def set_margins(doc: Document):
    for section in doc.sections:
        section.top_margin    = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin   = Inches(1.25)
        section.right_margin  = Inches(1.25)


def add_centered(doc, text, size=11, bold=False, italic=False, color=None, space_before=0, space_after=6):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after  = Pt(space_after)
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def add_section_heading(doc, number, title):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(16)
    para.paragraph_format.space_after  = Pt(4)
    run = para.add_run(f"{number}.  {title.upper()}")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = DARK_BLUE

    # Underline via a thin divider paragraph
    div = doc.add_paragraph()
    div.paragraph_format.space_before = Pt(0)
    div.paragraph_format.space_after  = Pt(8)
    div_run = div.add_run("─" * 85)
    div_run.font.size = Pt(8)
    div_run.font.color.rgb = GREY


def _add_formatted_run(para, text: str):
    """Render inline **bold** and *italic* markdown into Word runs."""
    para.paragraph_format.space_after = Pt(4)
    pattern = re.compile(r"(\*{1,3})(.*?)\1")
    last = 0
    for match in pattern.finditer(text):
        if match.start() > last:
            r = para.add_run(text[last:match.start()])
            r.font.size = Pt(10.5)
        markers, content = match.group(1), match.group(2)
        r = para.add_run(content)
        r.font.size = Pt(10.5)
        r.bold   = len(markers) >= 2
        r.italic = len(markers) in (1, 3)
        last = match.end()
    if last < len(text):
        r = para.add_run(text[last:])
        r.font.size = Pt(10.5)


def render_section(doc: Document, text: str):
    """Render LLM section text into the Word document."""
    # Pre-clean
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"^>\s*", "", text, flags=re.MULTILINE)
    text = re.sub(r"`{1,3}(.*?)`{1,3}", r"\1", text)
    text = re.sub(r"(?<!\*)\*(?!\*)", "", text)   # lone asterisks
    text = re.sub(r"\n{3,}", "\n\n", text)

    lines = text.splitlines()
    # Strip lines that are just a label with no value (e.g. "PREMIUM IMPACT:")
    lines = [l for l in lines if not re.match(r'^[A-Z][A-Z\s]{2,}:\s*$', l.strip())]
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip():
            i += 1
            continue

        # Markdown table block → Word table
        if line.strip().startswith("|") and line.strip().endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            data_rows = [r for r in table_lines if not re.match(r"^\|[-|:\s]+\|$", r.strip())]
            if data_rows:
                rows_parsed = [[c.strip() for c in r.strip().strip("|").split("|")] for r in data_rows]
                col_count = max(len(r) for r in rows_parsed)
                tbl = doc.add_table(rows=len(rows_parsed), cols=col_count)
                tbl.style = "Table Grid"
                for ri, row_data in enumerate(rows_parsed):
                    for ci, cell_text in enumerate(row_data):
                        if ci < col_count:
                            cell = tbl.rows[ri].cells[ci]
                            cell.text = cell_text
                            if ri == 0:
                                for run in cell.paragraphs[0].runs:
                                    run.bold = True
                doc.add_paragraph()
            continue

        # Markdown heading → bold paragraph
        if re.match(r"^#{1,6}\s+", line):
            heading_text = re.sub(r"^#{1,6}\s+", "", line).strip()
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(10)
            para.paragraph_format.space_after  = Pt(3)
            run = para.add_run(heading_text)
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = MID_BLUE

        # Bullet / numbered list
        elif re.match(r"^(\s*[-*•]|\s*\d+[.)]\s)", line):
            bullet_text = re.sub(r"^\s*[-*•]\s*", "", line)
            bullet_text = re.sub(r"^\s*\d+[.)]\s*", "", bullet_text)
            _add_formatted_run(doc.add_paragraph(style="List Bullet"), bullet_text.strip())

        # Body paragraph
        else:
            _add_formatted_run(doc.add_paragraph(), line.strip())

        i += 1


# ── Main builder ───────────────────────────────────────────────────────────────

def generate_document(
    request: str,
    policy_details: str,
    tasks: list[TaskItem],
    reflection: ReflectionNote,
) -> str:

    doc = Document()
    set_margins(doc)

    # ── COVER PAGE ─────────────────────────────────────────────────────────────
    for _ in range(4):
        doc.add_paragraph()

    add_centered(doc, "INSURANCE POLICY DOCUMENT", size=22, bold=True, color=DARK_BLUE, space_after=4)
    add_centered(doc, "━" * 42, size=12, color=DARK_BLUE, space_after=12)
    add_centered(doc, INSURER_NAME, size=12, bold=True, space_after=4)
    add_centered(doc, INSURER_ADDRESS, size=10, color=GREY, space_after=2)
    add_centered(doc, f"Tel: {INSURER_PHONE}  |  {INSURER_WEB}", size=10, color=GREY, space_after=12)
    add_centered(doc, "━" * 42, size=12, color=DARK_BLUE, space_after=12)
    add_centered(doc, f"Date Issued:  {datetime.now().strftime('%B %d, %Y')}", size=11, space_after=4)

    doc.add_page_break()

    # ── 9 POLICY SECTIONS ──────────────────────────────────────────────────────
    section_number = 1
    for task in tasks:
        if task.status == "completed" and task.result:
            add_section_heading(doc, section_number, task.task_name)
            render_section(doc, task.result)
            doc.add_page_break()
            section_number += 1

    # ── SAVE ───────────────────────────────────────────────────────────────────
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    unique_id = uuid.uuid4().hex[:8]
    filename  = f"insurance_policy_{timestamp}_{unique_id}.docx"
    filepath  = OUTPUT_DIR / filename

    try:
        doc.save(str(filepath))
    except IOError as e:
        raise DocumentError(f"Failed to save document to {filepath}: {e}") from e

    logger.info("Document saved: %s", filepath)
    return str(filepath)
